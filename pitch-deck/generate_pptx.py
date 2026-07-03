#!/usr/bin/env python3
"""Dead Man's Vault pitch deck — version 2 (card grids)."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

BG = RGBColor(0x07, 0x09, 0x0F)
SURFACE = RGBColor(0x0F, 0x15, 0x21)
ACCENT = RGBColor(0x00, 0xFF, 0xA3)
PURPLE = RGBColor(0x99, 0x45, 0xFF)
WARNING = RGBColor(0xF5, 0x9E, 0x0B)
CRITICAL = RGBColor(0xEF, 0x44, 0x44)
DARK_RED = RGBColor(0x99, 0x1B, 0x1B)
ORANGE = RGBColor(0xF9, 0x73, 0x16)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0xAA, 0xAA, 0xAA)
DIM = RGBColor(0x66, 0x66, 0x66)
BLUE = RGBColor(0x60, 0xA5, 0xFA)
TEAL = RGBColor(0x14, 0xF1, 0x95)

FONT = 'Calibri'

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW = 13.333
SH = 7.5
MARGIN = 1.0


def make_slide():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    return s


def text(slide, l, t, w, h, txt, sz=18, clr=WHITE, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    box.text_frame.word_wrap = True
    p = box.text_frame.paragraphs[0]
    p.text = txt
    p.alignment = align
    r = p.runs[0]
    r.font.size = Pt(sz)
    r.font.color.rgb = clr
    r.font.bold = bold
    r.font.name = FONT
    return box


def rect(slide, l, t, w, h, fill=SURFACE):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.adjustments[0] = 0.04
    return s


def bar(slide, l, t, w, color):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(0.05))
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()


def dot(slide, l, t, color, size=0.18):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(l), Inches(t), Inches(size), Inches(size))
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()


def footer(slide, n, total=8):
    text(slide, MARGIN, SH - 0.5, 3, 0.3, "Dead Man's Vault", 10, DIM, bold=True)
    text(slide, SW - MARGIN - 1.5, SH - 0.5, 1.5, 0.3, f"{n} / {total}", 10, DIM, align=PP_ALIGN.RIGHT)


# ── SLIDE 1: Title ──
s = make_slide()
rect(s, MARGIN, 1.0, 4.8, 0.4, RGBColor(0x1A, 0x0F, 0x2E))
text(s, MARGIN, 1.0, 4.8, 0.4, "MONOLITH  —  SOLANA MOBILE HACKATHON 2026", 11, PURPLE, bold=True, align=PP_ALIGN.CENTER)
text(s, MARGIN, 2.0, 10, 0.9, "Dead Man's Vault", 56, WHITE, bold=True)
text(s, MARGIN, 3.2, 10, 0.5, "An autonomous crypto inheritance protocol for Solana Seeker.", 22, GRAY)
text(s, MARGIN, 3.8, 10, 0.5, "Your crypto should outlive you.", 22, ACCENT, bold=True)
text(s, MARGIN, 5.2, 10, 0.3, "by Palatine Arc   ·   dmv.palatinearc.com   ·   github.com/PalatineArcOrg/DMV", 13, DIM)
footer(s, 1)


# ── SLIDE 2: Problem ──
s = make_slide()
text(s, MARGIN, 0.7, 3, 0.3, "THE PROBLEM", 12, ACCENT, bold=True)
text(s, MARGIN, 1.2, 10, 0.8, "Crypto doesn't have a \"next of kin\" form.", 42, WHITE, bold=True)
text(s, MARGIN, 2.3, 9, 0.7, "Millions in crypto assets are lost every year because there's no decentralized way to pass them on. Traditional inheritance systems don't work for self-custodied wallets.", 16, GRAY)

stats = [
    ("$140B+", ACCENT, "Estimated value of inaccessible\nBitcoin alone. Keys lost forever."),
    ("4M+", WARNING, "Bitcoin wallets with no activity\nin 5+ years. Many owners gone."),
    ("0", CRITICAL, "Decentralized inheritance solutions\non Solana Mobile. Until now."),
]
for i, (num, color, desc) in enumerate(stats):
    x = MARGIN + i * 3.8
    rect(s, x, 3.8, 3.5, 2.4)
    text(s, x + 0.35, 4.0, 2.8, 0.7, num, 40, color, bold=True)
    text(s, x + 0.35, 4.8, 2.8, 1.0, desc, 13, GRAY)
footer(s, 2)


# ── SLIDE 3: How It Works ──
s = make_slide()
text(s, MARGIN, 0.7, 3, 0.3, "HOW IT WORKS", 12, ACCENT, bold=True)
text(s, MARGIN, 1.2, 10, 0.8, "Five stages. One tap to stay alive.", 42, WHITE, bold=True)
text(s, MARGIN, 2.2, 9, 0.5, "Set up your vault once. Confirm a heartbeat periodically. If you stop, the protocol takes over.", 16, GRAY)

stages = [
    ("0", "Normal", "Heartbeat confirmed.\nAll clear. Assets safe.", "Weekly / monthly", ACCENT),
    ("1", "Reminder", "Overdue. Notifications\nsent to owner.", "3 days", WARNING),
    ("2", "Alert", "No response. Frequency\nincreases.", "7 days", ORANGE),
    ("3", "Warning", "Final countdown.\nHourly notifications.", "7 days", CRITICAL),
    ("4", "Execution", "Irreversible. Assets\ndistributed on-chain.", "Autonomous", DARK_RED),
]

cw = 2.1
gap = 0.15
total_w = 5 * cw + 4 * gap
x0 = (SW - total_w) / 2

for i, (num, name, desc, timing, clr) in enumerate(stages):
    x = x0 + i * (cw + gap)
    y = 3.1
    ch = 3.5

    rect(s, x, y, cw, ch)
    bar(s, x, y, cw, clr)
    text(s, x + 0.2, y + 0.25, 0.4, 0.3, num, 16, clr, bold=True)
    text(s, x + 0.2, y + 0.65, cw - 0.4, 0.3, name, 14, WHITE, bold=True)
    text(s, x + 0.2, y + 1.1, cw - 0.4, 1.0, desc, 11, GRAY)
    text(s, x + 0.2, y + 2.9, cw - 0.4, 0.3, timing, 9, DIM, bold=True)
footer(s, 3)


# ── SLIDE 4: Features ──
s = make_slide()
text(s, MARGIN, 0.7, 3, 0.3, "FEATURES", 12, ACCENT, bold=True)
text(s, MARGIN, 1.2, 10, 0.8, "Built for real-world crypto inheritance.", 42, WHITE, bold=True)

features = [
    ("TEE-Secured Agent Key", "Dedicated agent key in Seeker's\nhardware enclave. Handles heartbeats\nand execution autonomously."),
    ("Vault PDA Storage", "Deposit SOL and SPL tokens into\nyour vault PDA. Withdraw anytime.\nAgent distributes at Stage 4."),
    ("Crash-Proof Execution", "Every step checkpointed to SQLite.\nApp crashes, device restarts —\nexecution resumes where it left off."),
    ("Live Portfolio Tracking", "Token balances via Helius DAS,\ndual-oracle USD pricing (Pyth +\nJupiter), DeFi detection (10 protocols)."),
    ("On-Chain Whitelist", "Program only sends to pre-approved\nwallets. Up to 20 beneficiaries with\nbasis-point share allocation."),
    ("Full Transparency", "Every heartbeat, distribution, state\nchange on-chain. Execution logs link\ndirectly to Solana Explorer."),
]

cw = 3.5
gap = 0.3
for i, (title, desc) in enumerate(features):
    col = i % 3
    row = i // 3
    x = MARGIN + col * (cw + gap)
    y = 2.5 + row * 2.5
    rect(s, x, y, cw, 2.2)
    text(s, x + 0.3, y + 0.25, cw - 0.6, 0.3, title, 15, WHITE, bold=True)
    text(s, x + 0.3, y + 0.7, cw - 0.6, 1.3, desc, 12, GRAY)
footer(s, 4)


# ── SLIDE 5: Architecture ──
s = make_slide()
text(s, MARGIN, 0.7, 3, 0.3, "ARCHITECTURE", 12, ACCENT, bold=True)
text(s, MARGIN, 1.2, 10, 0.8, "No backend. No intermediaries.", 42, WHITE, bold=True)
text(s, MARGIN, 2.3, 9, 0.5, "Everything runs on-device or on-chain. The mobile app orchestrates, the Anchor program enforces, and the TEE secures.", 16, GRAY)

layers = [
    ("Anchor Program (Rust)", "13 instructions · On-chain enforcement", ACCENT),
    ("React Native (Expo SDK 52)", "Orchestration, scanning, escalation", PURPLE),
    ("Zustand + SQLite", "State management + crash recovery", WARNING),
    ("Solana Mobile MWA", "Owner authorization via Seed Vault", TEAL),
    ("Seeker TEE", "Agent key for autonomous operations", CRITICAL),
    ("Helius + Pyth + Jupiter", "Portfolio data, price oracles, priority fees", BLUE),
]

cw = 5.3
gap = 0.3
for i, (name, desc, clr) in enumerate(layers):
    col = i % 2
    row = i // 2
    x = MARGIN + col * (cw + gap)
    y = 3.3 + row * 1.2
    rect(s, x, y, cw, 0.95)
    dot(s, x + 0.3, y + 0.38, clr)
    text(s, x + 0.65, y + 0.15, cw - 1, 0.3, name, 14, WHITE, bold=True)
    text(s, x + 0.65, y + 0.5, cw - 1, 0.3, desc, 11, DIM)
footer(s, 5)


# ── SLIDE 6: On-Chain Program ──
s = make_slide()
text(s, MARGIN, 0.7, 4, 0.3, "ON-CHAIN PROGRAM", 12, ACCENT, bold=True)
text(s, MARGIN, 1.2, 10, 0.8, "Minimal program. Maximum security.", 42, WHITE, bold=True)
text(s, MARGIN, 2.3, 9, 0.5, "The program only enforces who can transfer, where transfers go, and when they're allowed. All complex logic lives in the mobile app.", 16, GRAY)

points = [
    ("13 Instructions", "Initialize, update, heartbeat, SOL/SPL distribution, record execution, rotate agent, revoke, withdraw, close."),
    ("19 Error Codes", "Interval validation, share allocation, signer auth, vault state guards, beneficiary whitelist, immutability, lifecycle."),
    ("34 Tests Passing", "Happy paths, error cases, security guards, SOL/SPL distribution, vault withdraw, close, double-execution prevention."),
    ("Mutable or Immutable", "Choose whether your vault can be updated/revoked, or permanent. Agent rotation always allowed."),
]

for i, (title, desc) in enumerate(points):
    y = 3.3 + i * 1.0
    dot(s, MARGIN + 0.05, y + 0.08, ACCENT, 0.15)
    text(s, MARGIN + 0.4, y, 4, 0.3, title, 15, WHITE, bold=True)
    text(s, MARGIN + 0.4, y + 0.35, 10, 0.4, desc, 12, GRAY)

rect(s, MARGIN, 6.5, 8.5, 0.45)
text(s, MARGIN + 0.25, 6.5, 8, 0.45, "Program ID:  GXCu5964mvgAJDWmcMriZpzU3vDVqPzjYCM1sxCnsoEb  (Devnet)", 12, ACCENT)
footer(s, 6)


# ── SLIDE 7: Differentiators ──
s = make_slide()
text(s, MARGIN, 0.7, 5, 0.3, "WHY DEAD MAN'S VAULT", 12, ACCENT, bold=True)
text(s, MARGIN, 1.2, 10, 0.8, "What makes us different.", 42, WHITE, bold=True)

diffs = [
    ("Truly Serverless", "No backend, no cloud functions, no intermediaries. Everything on-device or on-chain."),
    ("Built for Seeker", "First-class Solana Mobile integration. Agent key in TEE hardware enclave. Owner signs via Seed Vault MWA."),
    ("Owner Supremacy", "Owner can always override, withdraw, revoke, or rotate. Program enforces rules even if device is compromised."),
    ("Graduated Escalation", "17+ days of warnings before execution. Any heartbeat at Stages 1-3 resets. Stage 4 is the only irreversible action."),
]

for i, (title, desc) in enumerate(diffs):
    y = 2.5 + i * 1.1
    dot(s, MARGIN + 0.05, y + 0.08, ACCENT, 0.15)
    text(s, MARGIN + 0.4, y, 4, 0.3, title, 15, WHITE, bold=True)
    text(s, MARGIN + 0.4, y + 0.35, 10, 0.4, desc, 12, GRAY)

text(s, MARGIN, 6.2, 3, 0.3, "TECH STACK", 12, ACCENT, bold=True)

techs = ["Anchor 0.32.1", "Solana CLI 3.0.15", "Rust 1.93", "Expo SDK 52", "React Native 0.76.9", "TypeScript 5.x", "Zustand 5"]
px = MARGIN
for t in techs:
    w = len(t) * 0.095 + 0.45
    rect(s, px, 6.55, w, 0.35)
    text(s, px, 6.55, w, 0.35, t, 11, GRAY, align=PP_ALIGN.CENTER)
    px += w + 0.15
footer(s, 7)


# ── SLIDE 8: CTA ──
s = make_slide()
text(s, MARGIN, 0.7, 3, 0.3, "GET STARTED", 12, ACCENT, bold=True)
text(s, MARGIN, 1.2, 10, 0.8, "Download Dead Man's Vault", 42, WHITE, bold=True)
text(s, MARGIN, 2.3, 9, 0.5, "Currently live on Solana Devnet. Mainnet launch with full app store distribution coming soon.", 16, GRAY)

cards = [
    ("Devnet APK", "Test the full flow on Solana\nDevnet. Requires a Solana\nMobile wallet.", "Available Now", ACCENT),
    ("GitHub", "Open-source. MIT licensed.\nFull Anchor program +\nReact Native app.", "Open Source", ACCENT),
    ("Mainnet", "Production release with\nsecurity audit and Google\nPlay distribution.", "Coming Soon", PURPLE),
]

cw = 3.5
gap = 0.3
for i, (title, desc, badge_txt, badge_clr) in enumerate(cards):
    x = MARGIN + i * (cw + gap)
    y = 3.3
    rect(s, x, y, cw, 2.8)
    text(s, x, y + 0.3, cw, 0.4, title, 18, WHITE, bold=True, align=PP_ALIGN.CENTER)
    text(s, x + 0.3, y + 0.9, cw - 0.6, 1.0, desc, 12, GRAY, align=PP_ALIGN.CENTER)
    bw = len(badge_txt) * 0.09 + 0.5
    bx = x + (cw - bw) / 2
    rect(s, bx, y + 2.2, bw, 0.33, SURFACE)
    text(s, bx, y + 2.2, bw, 0.33, badge_txt, 10, badge_clr, bold=True, align=PP_ALIGN.CENTER)

text(s, 0, 6.6, SW, 0.3, "Web: dmv.palatinearc.com   ·   GitHub: github.com/PalatineArcOrg/DMV   ·   APK: GitHub Releases", 13, ACCENT, align=PP_ALIGN.CENTER)
footer(s, 8)

out = "/root/DMV/pitch-deck/Dead_Mans_Vault_Pitch_Deck.pptx"
prs.save(out)
print(f"Saved: {out}")
